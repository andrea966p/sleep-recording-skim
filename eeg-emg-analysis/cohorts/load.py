#!/usr/bin/env python3
"""Load a cohort manifest, and optionally left-join a husbandry spreadsheet.

    from cohorts.load import load_cohort
    cohort = load_cohort("HLB-fem-profsurgery-troubleshooting")
    cohort = load_cohort(path, husbandry="~/mice.xlsx")

THE JOIN
--------
The YAML manifest is the skeleton and is under version control. The husbandry
spreadsheet is the living record and is not. The spreadsheet WINS on any field
it supplies, because it is the one a human keeps up to date, and every value it
supplies is stamped with `source="husbandry-sheet:<filename>"` so a card can
always say where a number came from.

The loader reports which columns it recognised and which it ignored. It never
silently drops a column, because a typo in a header is otherwise invisible.

COLUMN SCHEMA
-------------
Header matching is case-insensitive and ignores spaces, underscores and hyphens,
so "Body weight at surgery", "body_weight_at_surgery" and "BodyWeightAtSurgery"
are the same column. Recognised names are listed in COLUMN_ALIASES below; add to
that dict rather than renaming anyone's spreadsheet.
"""
from __future__ import annotations

import sys
from pathlib import Path
from typing import Any

import yaml


from cohorts.schema import (  # noqa: E402
    Animal, Cohort, Intervention, RecordingRef, RigPlacement, Weighing, as_date,
)

COHORT_DIR = Path(__file__).resolve().parent


# --------------------------------------------------------------------------- #
# spreadsheet column schema
# --------------------------------------------------------------------------- #

def _norm(header: str) -> str:
    return "".join(ch for ch in str(header).lower() if ch.isalnum())


# canonical field -> the header spellings that mean it
COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "animal_id":            ("animalid", "animal", "mouseid", "mouse", "id", "eartag"),
    "project_id":           ("projectid", "project"),
    "batch_id":             ("batchid", "batch", "cohort", "cohortid"),
    "sex":                  ("sex", "gender"),
    "vendor":               ("vendor", "supplier", "company", "source"),
    "strain":               ("strain", "background", "genotype"),
    "arrival_date":         ("arrivaldate", "arrival", "datearrived", "receiveddate"),
    "age_at_arrival_weeks": ("ageatarrivalweeks", "ageatarrival", "agearrivalweeks",
                             "ageonarrival", "ageweeksatarrival"),
    "dob":                  ("dob", "dateofbirth", "birthdate", "birthday"),
    "surgery_date":         ("surgerydate", "surgery", "datesurgery", "implantdate",
                             "dateofsurgery"),
    "surgeon":              ("surgeon", "operator", "performedby"),
    "surgery_notes":        ("surgerynotes", "surgicalnotes", "notessurgery"),
    "tether_date":          ("tetherdate", "tethering", "tetheringdate", "datetethered",
                             "connectiondate"),
    "sacrifice_date":       ("sacrificedate", "sacrifice", "euthanasiadate", "datesacrificed"),
    "sampling_date":        ("samplingdate", "sampling", "collectiondate", "harvestdate"),
    "sampling_notes":       ("samplingnotes", "notessampling", "harvestnotes"),
    "notes":                ("notes", "comment", "comments", "remarks"),
    # weighings, each becomes a dated Weighing tied to its occasion
    "weight_arrival":       ("bodyweightatarrival", "weightatarrival", "arrivalweight",
                             "bwarrival"),
    "weight_surgery":       ("bodyweightatsurgery", "weightatsurgery", "surgeryweight",
                             "bwsurgery"),
    "weight_tether":        ("bodyweightattethering", "weightattethering", "tetherweight",
                             "bwtethering", "bodyweightattether"),
    "weight_sacrifice":     ("bodyweightatsacrifice", "weightatsacrifice", "sacrificeweight",
                             "bwsacrifice", "bodyweightatsampling"),
    # rig state
    "position":             ("position", "rigposition", "slot", "port"),
    "tether_arm":           ("tetherarm", "arm", "tethertype", "tether"),
}

_WEIGHT_OCCASIONS = {
    "weight_arrival":   ("arrival", "arrival_date"),
    "weight_surgery":   ("surgery", "surgery_date"),
    "weight_tether":    ("tethering", "tether_date"),
    "weight_sacrifice": ("sacrifice", "sacrifice_date"),
}

# Interventions are not fixed columns. Any header that starts with one of these
# is read as a dated intervention whose kind is the rest of the header.
_INTERVENTION_PREFIXES = ("intervention", "experiment", "treatment", "test")

_ALIAS_TO_FIELD = {alias: fieldname
                   for fieldname, aliases in COLUMN_ALIASES.items()
                   for alias in aliases}


# --------------------------------------------------------------------------- #
# manifest
# --------------------------------------------------------------------------- #

def load_manifest(path_or_id: str | Path) -> Cohort:
    """Load the YAML manifest. Accepts a path or a bare cohort id."""
    path = Path(path_or_id).expanduser()
    if not path.exists():
        candidate = COHORT_DIR / f"{path_or_id}.yaml"
        if candidate.exists():
            path = candidate
        else:
            raise SystemExit(
                f"no cohort manifest at {path} and no {candidate}\n"
                f"available: {', '.join(sorted(p.stem for p in COHORT_DIR.glob('*.yaml'))) or 'none'}")
    obj = yaml.safe_load(path.read_text()) or {}
    cohort = Cohort.from_obj(obj)
    cohort_path = path
    setattr(cohort, "manifest_path", str(cohort_path))
    return cohort


# --------------------------------------------------------------------------- #
# husbandry spreadsheet
# --------------------------------------------------------------------------- #

def read_husbandry(path: str | Path) -> tuple[list[dict[str, Any]], dict[str, list[str]]]:
    """Read a CSV or XLSX husbandry sheet.

    Returns (rows, report) where report has "recognised" and "ignored" header
    lists so the caller can print exactly what was understood.
    """
    path = Path(path).expanduser()
    if not path.exists():
        raise SystemExit(f"husbandry sheet not found: {path}")

    if path.suffix.lower() in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        ws = wb[wb.sheetnames[0]]
        rows_raw = list(ws.iter_rows(values_only=True))
        wb.close()
        if not rows_raw:
            return [], {"recognised": [], "ignored": []}
        headers = [str(h) if h is not None else "" for h in rows_raw[0]]
        records = [dict(zip(headers, r)) for r in rows_raw[1:]]
    else:
        import csv
        with path.open(newline="", encoding="utf-8-sig") as fh:
            reader = csv.DictReader(fh)
            headers = list(reader.fieldnames or [])
            records = list(reader)

    recognised, ignored = [], []
    for h in headers:
        if not h:
            continue
        n = _norm(h)
        if n in _ALIAS_TO_FIELD or any(n.startswith(p) for p in _INTERVENTION_PREFIXES):
            recognised.append(h)
        else:
            ignored.append(h)
    return records, {"recognised": recognised, "ignored": ignored}


def _apply_row(animal: Animal, row: dict[str, Any], stamp: str) -> None:
    """Overlay one spreadsheet row onto an Animal. The sheet wins."""
    scalars: dict[str, Any] = {}
    weights: dict[str, Any] = {}
    interventions: list[Intervention] = []
    position = tether_arm = None

    for header, value in row.items():
        if header is None or value in (None, ""):
            continue
        n = _norm(header)
        fieldname = _ALIAS_TO_FIELD.get(n)
        if fieldname is None:
            if any(n.startswith(p) for p in _INTERVENTION_PREFIXES):
                # header carries the kind, cell carries the date
                kind = str(header)
                for p in _INTERVENTION_PREFIXES:
                    if n.startswith(p):
                        kind = str(header)[len(p):].strip(" _-:") or str(header)
                        break
                try:
                    when = as_date(value)
                except ValueError:
                    continue
                interventions.append(Intervention(kind=kind, start=when, end=when,
                                                  source=stamp))
            continue
        if fieldname in _WEIGHT_OCCASIONS:
            weights[fieldname] = value
        elif fieldname == "position":
            position = str(value)
        elif fieldname == "tether_arm":
            tether_arm = str(value)
        else:
            scalars[fieldname] = value

    date_fields = {"arrival_date", "dob", "surgery_date", "tether_date",
                   "sacrifice_date", "sampling_date"}
    for fieldname, value in scalars.items():
        if fieldname == "animal_id":
            continue
        if fieldname in date_fields:
            value = as_date(value)
        elif fieldname == "age_at_arrival_weeks":
            value = int(round(float(value)))
        else:
            value = str(value).strip()
        setattr(animal, fieldname, value)
        animal.sources[fieldname] = stamp

    for key, raw in weights.items():
        occasion, date_field = _WEIGHT_OCCASIONS[key]
        try:
            grams = float(raw)
        except (TypeError, ValueError):
            continue
        when = getattr(animal, date_field, None)
        existing = next((w for w in animal.weighings if w.occasion == occasion), None)
        if existing is not None:
            existing.grams, existing.date, existing.source = grams, when, stamp
        else:
            animal.weighings.append(
                Weighing(date=when, grams=grams, occasion=occasion, source=stamp))

    for iv in interventions:
        if not any(x.kind == iv.kind and x.start == iv.start for x in animal.interventions):
            animal.interventions.append(iv)

    if position or tether_arm:
        when = animal.tether_date or animal.surgery_date
        existing = next((p for p in animal.placements
                         if p.from_date == when), None)
        if existing is None:
            animal.placements.append(RigPlacement(
                from_date=when, position=position or "",
                tether_arm=tether_arm or "", source=stamp))
        else:
            if position:
                existing.position = position
            if tether_arm:
                existing.tether_arm = tether_arm
            existing.source = stamp

    animal.weighings.sort(key=lambda w: (w.date is None, w.date))
    animal.interventions.sort(key=lambda i: (i.start is None, i.start))
    animal.placements.sort(key=lambda p: (p.from_date is None, p.from_date))


def join_husbandry(cohort: Cohort, path: str | Path, *, verbose: bool = True) -> Cohort:
    """Left-join a husbandry sheet onto the cohort, matching on animal_id."""
    rows, report = read_husbandry(path)
    stamp = f"husbandry-sheet:{Path(path).name}"

    matched, unmatched = 0, []
    for row in rows:
        aid = None
        for header, value in row.items():
            if header and _ALIAS_TO_FIELD.get(_norm(header)) == "animal_id":
                aid = str(value).strip() if value is not None else None
                break
        if not aid:
            continue
        animal = cohort.by_id(aid)
        if animal is None:
            unmatched.append(aid)
            continue
        _apply_row(animal, row, stamp)
        matched += 1

    if verbose:
        print(f"husbandry sheet: {Path(path).name}")
        print(f"  columns recognised ({len(report['recognised'])}): "
              f"{', '.join(report['recognised']) or 'none'}")
        print(f"  columns IGNORED   ({len(report['ignored'])}): "
              f"{', '.join(report['ignored']) or 'none'}")
        print(f"  rows joined: {matched}")
        if unmatched:
            print(f"  ⚠️  rows with an animal_id NOT in the manifest: "
                  f"{', '.join(unmatched)}")
    return cohort


# --------------------------------------------------------------------------- #
# public entry point
# --------------------------------------------------------------------------- #

def load_cohort(path_or_id: str | Path,
                husbandry: str | Path | None = None,
                *, strict: bool = True, verbose: bool = True) -> Cohort:
    """Load a manifest, optionally overlay a husbandry sheet, then validate."""
    cohort = load_manifest(path_or_id)
    if husbandry:
        join_husbandry(cohort, husbandry, verbose=verbose)

    problems = cohort.validate()
    if problems:
        header = f"cohort {cohort.cohort_id}: {len(problems)} problem(s)"
        body = "\n".join(f"  - {p}" for p in problems)
        if strict:
            raise SystemExit(f"{header}\n{body}\n"
                             "Dates that cannot be true are refused rather than drawn.")
        print(f"⚠️  {header}\n{body}", file=sys.stderr)
    return cohort


def _main(argv: list[str]) -> int:
    import argparse
    ap = argparse.ArgumentParser(description="Inspect a cohort manifest.")
    ap.add_argument("cohort", help="cohort id or path to a manifest YAML")
    ap.add_argument("--husbandry", help="CSV or XLSX husbandry sheet to join")
    ap.add_argument("--no-strict", action="store_true",
                    help="report validation problems instead of refusing")
    args = ap.parse_args(argv)

    cohort = load_cohort(args.cohort, args.husbandry, strict=not args.no_strict)
    print(f"\ncohort   {cohort.cohort_id}")
    print(f"title    {cohort.title}")
    print(f"roles    {', '.join(cohort.roles)}   criterion: {cohort.criterion}")
    print(f"animals  {len(cohort.animals)}")
    print(f"{'animal':<14} {'sex':<4} {'pos':<5} {'arm':<4} "
          f"{'surgery':<11} {'tether':<11} {'recordings':<11} status")
    print("-" * 84)
    for a in cohort.animals:
        last = a.placements[-1] if a.placements else None
        print(f"{a.animal_id:<14} {a.sex or '-':<4} "
              f"{(last.position if last else '-'):<5} "
              f"{(last.tether_arm if last else '-'):<4} "
              f"{(a.surgery_date.isoformat() if a.surgery_date else '-'):<11} "
              f"{(a.tether_date.isoformat() if a.tether_date else '-'):<11} "
              f"{len(a.recordings):<11} "
              f"{'no end date recorded' if a.is_alive else 'sacrificed ' + a.sacrifice_date.isoformat()}")
    return 0


if __name__ == "__main__":
    raise SystemExit(_main(sys.argv[1:]))
